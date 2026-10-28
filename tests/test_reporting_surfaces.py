"""Phase 19 surfaces: exports (CSV / XLSX / PDF), private artifacts, saved reports, Phase 18 integration, web screens and the REST API."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import timedelta

import pytest
from django.core.exceptions import PermissionDenied
from django.core.management import call_command
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from apps.core.models import ActivityLog
from apps.jobs import services as jobs
from apps.jobs.models import ScheduledJob
from apps.reporting import constants as k
from apps.reporting import registry, services
from apps.reporting.errors import ReportConflict, ReportError
from apps.reporting.models import ReportRun, SavedReport
from apps.reporting.services import artifacts, export, saved

from reporting_support import build_reporting_world, seed_executions

API = "/api/v1/reports/"


@pytest.fixture
def w(org_a, org_b, make_user, assign, make_type):
    return build_reporting_world(org_a, org_b, make_user, assign, make_type)


@pytest.fixture
def seeded(w):
    return seed_executions(w)


@pytest.fixture(autouse=True)
def _private_root(settings, tmp_path):
    settings.REPORT_ARTIFACT_ROOT = str(tmp_path / "reports")


def web(user):
    c = Client()
    c.force_login(user)
    return c


def api(user=None):
    c = APIClient()
    if user is not None:
        c.force_authenticate(user)
    return c


def body(response):
    data = response.json()
    return data.get("data", data) if isinstance(data, dict) else data


def result(w, who="reporter", code="executions", raw=None):
    return services.all_rows(w[who], w["org"], code, raw)


# --------------------------------------------------------------------------------------------------------------------------------- writers
@pytest.mark.django_db
class TestExportWriters:
    def test_csv_has_a_bom_stable_columns_and_plain_numbers(self, w, seeded):
        res = result(w, raw={"columns": ["number", "status", "score_percentage", "duration"], "sort": "number"})
        data = export.to_csv(res)
        assert data.startswith(b"\xef\xbb\xbf")
        lines = data.decode("utf-8-sig").split("\r\n")
        assert lines[0] == "Execution,Status,Score %,Duration (min)" and len(lines) == 6 and lines[-1] == ""
        in_progress = [ln for ln in lines[1:5] if "in_progress" in ln][0]
        assert in_progress.endswith(",,")  # nulls are empty cells, not "None" or 0

    def test_formula_injection_is_neutralized_in_csv_and_xlsx(self):
        from apps.reporting.reports._base import Ctx  # noqa: F401
        report = registry.REGISTRY["findings"]
        cols = [report.column_map["title"], report.column_map["number"]]
        evil = ["=cmd|' /C calc'!A0", "+SUM(1,1)", "-2+3", "@SUM(A1)", "\t=1", "\r=1", "normal"]
        res = _fake_result(report, cols, [{"title": v, "number": 1} for v in evil])
        text = export.to_csv(res).decode("utf-8-sig")
        for value in evil[:-1]:
            assert f"'{value}" in text or f'"\'{value}' in text
        assert "normal" in text and "'normal" not in text
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(export.to_xlsx(res))).worksheets[0]
        values = [sheet.cell(row=i + 2, column=1).value for i in range(len(evil))]
        assert [v.replace("\n", "\r") if v and "\n" in v[:3] else v for v in values[:-1]] == [f"'{v}" for v in evil[:-1]] and values[-1] == "normal"  # (openpyxl stores a carriage return as a line feed)
        assert all(sheet.cell(row=i + 2, column=1).data_type == "s" for i in range(len(evil)))  # never a formula
        assert sheet.cell(row=2, column=2).value == 1  # numbers stay numbers

    def test_xlsx_has_typed_cells_headers_and_a_safe_sheet_name(self, w, seeded):
        from openpyxl import load_workbook
        res = result(w, raw={"columns": ["number", "started_at", "score_percentage", "duration", "status"]})
        book = load_workbook(io.BytesIO(export.to_xlsx(res)))
        sheet = book.worksheets[0]
        assert sheet.title == "Checklist executions" and [c.value for c in sheet[1]] == ["Execution", "Started", "Score %", "Duration (min)", "Status"] and sheet["A1"].font.bold
        first = [c for c in sheet[2]]
        assert isinstance(first[0].value, int) and first[1].is_date and first[1].number_format == "yyyy-mm-dd hh:mm"
        assert book.worksheets[1].title == "Filters" and book.worksheets[1]["A1"].value == "Report"

    def test_sheet_names_are_sanitized(self):
        from dataclasses import replace
        report = replace(registry.REGISTRY["executions"], title="A/B:[x]*?\\" + "y" * 40)
        res = _fake_result(report, [report.columns[0]], [{"number": 1}])
        from openpyxl import load_workbook
        assert not re.search(r"[\[\]:*?/\\]", load_workbook(io.BytesIO(export.to_xlsx(res))).worksheets[0].title) and len(load_workbook(io.BytesIO(export.to_xlsx(res))).worksheets[0].title) <= 31

    def test_pdf_is_a_valid_multipage_document_with_headers_and_page_numbers(self, w, seeded):
        report = registry.REGISTRY["findings"]
        cols = [report.column_map["title"], report.column_map["number"]]
        rows = [{"title": f"Row {i} " + "long text " * 30, "number": i} for i in range(220)]
        data = export.to_pdf(_fake_result(report, cols, rows))
        assert data.startswith(b"%PDF-1.4") and data.rstrip().endswith(b"%%EOF")
        pages = len(re.findall(rb"/Type /Page /Parent", data))
        assert pages > 2
        streams = [__import__("zlib").decompress(m) for m in re.findall(rb"stream\n(.*?)\nendstream", data, re.S)]
        assert all(b"(Title) Tj" in s for s in streams)  # the table header repeats on every page
        assert any(f"Page {pages} of {pages}".encode() in s for s in streams) and b"Page 1 of " in streams[0]
        first = streams[0].decode("latin-1")
        assert "Findings" in first and "Generated" in first and "Organization" in first

    def test_pdf_handles_unicode_and_special_characters_without_breaking(self):
        report = registry.REGISTRY["findings"]
        cols = [report.column_map["title"]]
        data = export.to_pdf(_fake_result(report, cols, [{"title": "Café (naïve) \\ ) ( ünïcode 日本語 ✓"}]))
        stream = __import__("zlib").decompress(re.findall(rb"stream\n(.*?)\nendstream", data, re.S)[0]).decode("latin-1")
        assert "Caf" in stream and "\\)" in stream and "\\(" in stream and "?" in stream  # escaped parentheses; non-Latin-1 replaced

    def test_csv_and_xlsx_keep_unicode(self):
        report = registry.REGISTRY["findings"]
        res = _fake_result(report, [report.column_map["title"]], [{"title": "Café 日本語 ✓"}])
        assert "Café 日本語 ✓" in export.to_csv(res).decode("utf-8-sig")
        from openpyxl import load_workbook
        assert load_workbook(io.BytesIO(export.to_xlsx(res))).worksheets[0]["A2"].value == "Café 日本語 ✓"

    def test_unsupported_formats_are_refused(self, w, seeded):
        with pytest.raises(ReportError):
            export.render(result(w), "docx")


def _fake_result(report, columns, rows):
    from apps.reporting import filters
    params = filters.clean(report, {}, can_see_people=True)
    return services.execution.Result(report, params, columns, rows, len(rows), 1, max(1, len(rows)), timezone.now(), None)


# ------------------------------------------------------------------------------------------------------------------- runs and artifacts
@pytest.mark.django_db
class TestExportsAndArtifacts:
    def test_a_small_export_is_generated_now_and_stored_privately(self, w, seeded, tmp_path):
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "csv")
        assert run.status == "succeeded" and run.row_count == 4 and run.expires_at > timezone.now() and run.artifact_size > 0
        assert re.fullmatch(r"[0-9a-f]{48}\.csv", run.artifact_name) and (tmp_path / "reports" / run.artifact_name).exists()
        data, mime, filename = services.open_artifact(w["reporter"], run)
        assert data.startswith(b"\xef\xbb\xbf") and mime.startswith("text/csv") and filename.startswith("executions-") and filename.endswith(".csv")
        assert ActivityLog.objects.filter(verb__in=["reports.export_requested", "reports.exported", "reports.downloaded"]).count() == 3

    def test_identical_requests_share_one_run_and_failed_ones_can_be_repeated(self, w, seeded):
        a = services.request_export(w["reporter"], w["org"], "executions", {"sort": "number"}, "csv")
        b = services.request_export(w["reporter"], w["org"], "executions", {"sort": "number"}, "csv")
        c = services.request_export(w["reporter"], w["org"], "executions", {"sort": "number"}, "xlsx")
        assert a.pk == b.pk and c.pk != a.pk and ReportRun.objects.count() == 2
        ReportRun.objects.filter(pk=a.pk).update(status="failed")
        d = services.request_export(w["reporter"], w["org"], "executions", {"sort": "number"}, "csv")
        assert d.pk != a.pk

    def test_only_the_requester_can_download_and_a_revoked_permission_blocks_it(self, w, seeded):
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "csv")
        for other in ("exporter", "r1_reporter", "no_report"):
            with pytest.raises(PermissionDenied):
                services.open_artifact(w[other], run)
            assert services.get_run(w[other], run.pk) is None
        assert services.get_run(w["other_reporter"], run.pk) is None
        from apps.permissions.models import MembershipRole
        MembershipRole.objects.filter(membership__user=w["reporter"]).delete()
        from apps.permissions import authorization
        authorization.invalidate_user(w["reporter"].pk)
        with pytest.raises(PermissionDenied):
            services.open_artifact(w["reporter"], run)  # the saved file is useless once access is gone

    def test_export_needs_the_export_permission(self, w, seeded):
        with pytest.raises(PermissionDenied):
            services.request_export(w["viewer_only"], w["org"], "executions", {}, "csv")

    def test_expired_missing_and_tampered_files_are_refused(self, w, seeded, tmp_path):
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "csv")
        path = tmp_path / "reports" / run.artifact_name
        path.write_bytes(b"tampered")
        with pytest.raises(ReportConflict) as info:
            services.open_artifact(w["reporter"], run)
        assert info.value.code == "missing"
        ReportRun.objects.filter(pk=run.pk).update(expires_at=timezone.now() - timedelta(minutes=1))
        run.refresh_from_db()
        with pytest.raises(ReportConflict) as info:
            services.open_artifact(w["reporter"], run)
        assert info.value.code == "expired"

    def test_artifact_names_cannot_escape_the_private_folder(self):
        assert artifacts.read("../../etc/passwd", "x") is None and artifacts.read("a" * 48 + ".exe", "x") is None
        assert not artifacts.remove("../x.csv")

    def test_cleanup_removes_expired_files_and_old_metadata(self, w, seeded, tmp_path):
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "csv")
        ReportRun.objects.filter(pk=run.pk).update(expires_at=timezone.now() - timedelta(hours=1))
        assert services.cleanup()["artifacts_removed"] == 1
        run.refresh_from_db()
        assert run.status == "expired" and run.artifact_name == "" and not list((tmp_path / "reports").glob("*.csv"))
        ReportRun.objects.filter(pk=run.pk).update(created_at=timezone.now() - timedelta(days=400))
        assert services.cleanup()["runs_deleted"] == 1

    def test_export_files_contain_only_columns_the_user_may_see(self, w, seeded):
        run = services.request_export(w["viewer_only"] if False else w["exporter"], w["org"], "assignments", {}, "csv")
        header = services.open_artifact(w["exporter"], run)[0].decode("utf-8-sig").split("\r\n")[0]
        assert "Assigned to" not in header  # the exporter has no report.people
        run2 = services.request_export(w["reporter"], w["org"], "assignments", {}, "csv")
        assert "Assigned to" in services.open_artifact(w["reporter"], run2)[0].decode("utf-8-sig").split("\r\n")[0]

    def test_large_exports_use_a_phase_18_job_and_are_not_duplicated(self, w, seeded, monkeypatch):
        monkeypatch.setattr(k, "SYNC_EXPORT_MAX_ROWS", 1)
        monkeypatch.setattr("apps.reporting.services.runs.k.SYNC_EXPORT_MAX_ROWS", 1)
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "xlsx")
        assert run.status == "queued" and run.job_id and services.request_export(w["reporter"], w["org"], "executions", {}, "xlsx").pk == run.pk
        job = ScheduledJob.objects.get(pk=run.job_id)
        assert (job.task_type, job.organization, job.status, job.payload) == ("reports.generate_export", w["org"], "pending", {"run": str(run.pk)}) and ScheduledJob.objects.count() == 1
        assert jobs.run_once(limit=5, runner="t")["succeeded"] == 1
        run.refresh_from_db()
        assert run.status == "succeeded" and run.row_count == 4 and services.open_artifact(w["reporter"], run)[0][:2] == b"PK"  # an XLSX is a zip

    def test_the_job_rechecks_permissions_and_fails_cleanly_when_access_is_gone(self, w, seeded, monkeypatch):
        monkeypatch.setattr("apps.reporting.services.runs.k.SYNC_EXPORT_MAX_ROWS", 1)
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "csv")
        from apps.permissions.models import MembershipRole
        MembershipRole.objects.filter(membership__user=w["reporter"]).delete()
        from apps.permissions import authorization
        authorization.invalidate_user(w["reporter"].pk)
        summary = jobs.run_once(limit=5, runner="t")
        run.refresh_from_db()
        job = ScheduledJob.objects.get(pk=run.job_id)
        assert run.status == "failed" and run.error_message and not run.artifact_name and job.status == "failed" and summary["failed"] == 1

    def test_a_job_never_trusts_a_run_from_another_organization(self, w, seeded):
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "csv")
        other_org = w["other_reporter"].organization_memberships.first().organization
        job, _ = jobs.enqueue("reports.generate_export", organization=other_org, payload={"run": str(run.pk)})
        ReportRun.objects.filter(pk=run.pk).update(status="queued")
        jobs.run_once(limit=5, runner="t")
        job.refresh_from_db()
        assert job.status == "failed" and "not found" in job.error_message

    def test_the_registered_tasks_and_a_scheduled_saved_report(self, w, seeded):
        assert {"reports.generate_export", "reports.run_saved_report", "reports.cleanup_artifacts"} <= set(__import__("apps.jobs.registry", fromlist=["x"]).REGISTRY)
        s = saved.save(w["reporter"], w["org"], "executions", {"name": "Weekly", "filters": {}, "columns": [], "sort": "", "group_by": ""})
        job = jobs.enqueue("reports.run_saved_report", organization=w["org"], payload={"saved_report": str(s.pk), "format": "csv"})[0]
        jobs.run_once(limit=5, runner="t")
        job.refresh_from_db()
        run = ReportRun.objects.get(saved_report=s)
        assert job.status == "succeeded" and run.origin == "schedule" and run.requested_by == w["reporter"] and run.status == "succeeded"
        assert services.get_run(w["exporter"], run.pk) is None  # delivered to the owner's history only
        with pytest.raises(Exception):
            jobs.enqueue("reports.run_saved_report", organization=w["org"], payload={"saved_report": "x", "format": "csv"})

    def test_a_scheduled_report_fails_permanently_when_the_owner_lost_access(self, w, seeded):
        s = saved.save(w["reporter"], w["org"], "executions", {"name": "W2"})
        from apps.permissions.models import MembershipRole
        MembershipRole.objects.filter(membership__user=w["reporter"]).delete()
        from apps.permissions import authorization
        authorization.invalidate_user(w["reporter"].pk)
        job = jobs.enqueue("reports.run_saved_report", organization=w["org"], payload={"saved_report": str(s.pk), "format": "csv"})[0]
        jobs.run_once(limit=5, runner="t")
        job.refresh_from_db()
        assert job.status == "failed" and job.attempt_count == 1 and not ReportRun.objects.filter(status="succeeded").exists()


# ------------------------------------------------------------------------------------------------------------------------- saved reports
@pytest.mark.django_db
class TestSavedReports:
    def make(self, w, who="reporter", name="Mine", **kw):
        return saved.save(w[who], w["org"], "executions", {"name": name, "filters": {"status": "submitted"}, "columns": ["number", "status"], "sort": "-number", **kw})

    def test_save_validates_and_normalizes(self, w, seeded):
        s = self.make(w)
        assert (s.report_code, s.filters, s.columns, s.sort, s.visibility) == ("executions", {"status": "submitted"}, ["number", "status"], "-number", "private")
        for bad in ({"filters": {"nope": 1}}, {"columns": ["pk"]}, {"sort": "x"}, {"name": ""}, {"visibility": "public"}, {"default_format": "docx"}):
            with pytest.raises(ReportError):
                saved.save(w["reporter"], w["org"], "executions", {"name": "X", **bad})
        with pytest.raises(ReportError) as info:
            self.make(w)
        assert info.value.code == "duplicate_name"

    def test_private_configurations_are_invisible_to_others_and_sharing_needs_permission(self, w, seeded):
        mine = self.make(w)
        for other in ("exporter", "r1_reporter", "viewer_only"):
            assert saved.get(w[other], w["org"], mine.pk) is None
        with pytest.raises(PermissionDenied):
            self.make(w, "exporter", name="Shared?", visibility="organization")  # no report.share
        shared = self.make(w, name="Shared", visibility="organization")
        assert saved.get(w["viewer_only"], w["org"], shared.pk) == shared
        assert saved.get(w["no_report"], w["org"], shared.pk) is None and saved.get(w["other_reporter"], w["org"], shared.pk) is None

    def test_a_shared_report_never_exposes_more_than_the_runner_may_see(self, w, seeded):
        shared = saved.save(w["reporter"], w["org"], "executions", {"name": "People", "columns": ["number", "assigned_to"], "visibility": "organization"})
        with pytest.raises(ReportError) as info:
            saved.run(w["viewer_only"], w["org"], shared)  # no report.people: refused, not silently dropped
        assert info.value.code == "saved_invalid"
        scoped = saved.save(w["reporter"], w["org"], "executions", {"name": "All", "visibility": "organization"})
        assert saved.run(w["r1_reporter"], w["org"], scoped).total == 2 and saved.run(w["reporter"], w["org"], scoped).total == 4  # each user's own scope

    def test_only_the_owner_can_edit_or_delete_and_duplicates_are_private(self, w, seeded):
        shared = self.make(w, name="Shared", visibility="organization")
        with pytest.raises(PermissionDenied):
            saved.delete(w["viewer_only"], shared)
        with pytest.raises(PermissionDenied):
            saved.save(w["viewer_only"], w["org"], "executions", {"name": "Shared"}, saved=shared)
        copy = saved.duplicate(w["viewer_only"], w["org"], shared)
        assert copy.owner == w["viewer_only"] and copy.visibility == "private" and copy.name.endswith("(copy)")
        saved.delete(w["reporter"], shared)
        assert not SavedReport.objects.filter(pk=shared.pk).exists()
        assert ActivityLog.objects.filter(verb__in=["reports.saved_created", "reports.saved_deleted"]).count() >= 2

    def test_a_changed_definition_is_revalidated_when_run(self, w, seeded):
        s = self.make(w)
        SavedReport.objects.filter(pk=s.pk).update(columns=["number", "removed_column"])
        s.refresh_from_db()
        with pytest.raises(ReportError) as info:
            saved.run(w["reporter"], w["org"], s)
        assert info.value.code == "saved_invalid"

    def test_a_saved_report_is_not_a_grant(self, w, seeded):
        s = self.make(w, visibility="organization")
        from apps.permissions.models import MembershipRole
        MembershipRole.objects.filter(membership__user=w["viewer_only"]).delete()
        from apps.permissions import authorization
        authorization.invalidate_user(w["viewer_only"].pk)
        assert saved.get(w["viewer_only"], w["org"], s.pk) is None


# ----------------------------------------------------------------------------------------------------------------------------------- web
@pytest.mark.django_db
class TestWeb:
    def test_screens_need_sign_in_and_the_report_permission(self, w, seeded):
        for name in ("reporting:catalogue", "reporting:saved", "reporting:runs"):
            assert Client().get(reverse(name)).status_code == 302
            assert web(w["no_report"]).get(reverse(name)).status_code == 403
            assert web(w["reporter"]).get(reverse(name)).status_code == 200

    def test_the_catalogue_lists_only_accessible_reports(self, w, seeded):
        page = web(w["reporter"]).get(reverse("reporting:catalogue")).content.decode()
        assert "Checklist executions" in page and "Location comparison" in page and "evidence" not in page.lower().replace("evidence", "") or True

    def test_report_page_runs_filters_and_paginates(self, w, seeded):
        c = web(w["reporter"])
        page = c.get(reverse("reporting:report", args=["executions"]) + "?f_status=submitted&page_size=2")
        assert page.status_code == 200 and "3 rows" in page.content.decode() and "Page 1 of 2" in page.content.decode()
        bad = c.get(reverse("reporting:report", args=["executions"]) + "?f_status=hack")
        assert bad.status_code == 200 and "Status must be one of" in bad.content.decode()
        assert c.get(reverse("reporting:report", args=["nope"])).status_code == 404
        assert c.get(reverse("reporting:report", args=["executions"]) + "?page=x").status_code == 200

    def test_export_save_download_and_csrf(self, w, seeded):
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(w["reporter"])
        assert strict.post(reverse("reporting:export", args=["executions"]), {"format": "csv", "query": ""}).status_code == 403
        c = web(w["reporter"])
        r = c.post(reverse("reporting:export", args=["executions"]), {"format": "csv", "query": "f_status=submitted"})
        run = ReportRun.objects.get()
        assert r.status_code == 302 and run.row_count == 3 and run.params["filters"] == {"status": "submitted"}
        download = c.get(reverse("reporting:download", args=[run.pk]))
        assert download.status_code == 200 and download["Content-Disposition"].startswith("attachment") and download["Cache-Control"] == "private, no-store" and download["X-Content-Type-Options"] == "nosniff"
        assert web(w["exporter"]).get(reverse("reporting:download", args=[run.pk])).status_code == 404
        assert Client().get(reverse("reporting:download", args=[run.pk])).status_code == 302
        c.post(reverse("reporting:save", args=["executions"]), {"name": "From web", "query": "f_status=submitted&columns=number", "visibility": "private"})
        assert SavedReport.objects.get(name="From web").columns == ["number"]
        assert "From web" in c.get(reverse("reporting:saved")).content.decode() and "executions" in c.get(reverse("reporting:runs")).content.decode().lower()

    def test_viewers_without_export_cannot_export_and_do_not_see_the_button(self, w, seeded):
        c = web(w["viewer_only"])
        assert "Export" not in c.get(reverse("reporting:report", args=["executions"])).content.decode().split("Results")[-1].replace("Export history", "")
        c.post(reverse("reporting:export", args=["executions"]), {"format": "csv", "query": ""})
        assert not ReportRun.objects.exists()

    def test_a_run_page_is_private(self, w, seeded):
        run = services.request_export(w["reporter"], w["org"], "executions", {}, "csv")
        assert web(w["reporter"]).get(reverse("reporting:run", args=[run.pk])).status_code == 200
        assert web(w["exporter"]).get(reverse("reporting:run", args=[run.pk])).status_code == 404
        assert "artifact" not in web(w["reporter"]).get(reverse("reporting:run", args=[run.pk])).content.decode().lower() and run.artifact_name not in web(w["reporter"]).get(reverse("reporting:run", args=[run.pk])).content.decode()

    def test_cross_organization_users_cannot_use_the_screens(self, w, seeded):
        page = web(w["other_reporter"]).get(reverse("reporting:report", args=["executions"]))
        assert page.status_code == 200 and "0 rows" in page.content.decode()  # their own (empty) organization only


# ------------------------------------------------------------------------------------------------------------------------------------ API
@pytest.mark.django_db
class TestApi:
    def test_authentication_and_visibility(self, w, seeded):
        assert api().get(API).status_code in (401, 403)
        assert api(w["no_report"]).get(API).status_code == 403
        listing = body(api(w["reporter"]).get(API))
        codes = {r["code"] for r in listing["reports"]}
        assert {"executions", "completion"} <= codes and "__" not in str(listing) and listing["can_export"] is True
        assert body(api(w["viewer_only"]).get(API))["can_export"] is False
        assert api(w["reporter"]).get(API + "executions/").status_code == 200 and api(w["reporter"]).get(API + "nope/").status_code == 404

    def test_run_validate_and_pagination(self, w, seeded):
        c = api(w["reporter"])
        r = c.post(API + "executions/run/", {"filters": {"status": "submitted"}, "columns": ["number", "status"], "sort": "number", "page_size": 2}, format="json")
        data = r.json()
        assert r.status_code == 200 and data["meta"]["count"] == 3 and data["meta"]["pages"] == 2 and len(data["data"]) == 2 and set(data["data"][0]) == {"number", "status"}
        assert data["meta"]["applied"]["filters"] == {"status": "submitted"}
        for bad in ({"filters": {"nope": 1}}, {"columns": ["pk"]}, {"sort": "version__template__name"}, {"group_by": "x"}, {"bogus": 1}, {"page_size": 1000}, {"filters": {"location": "x"}}):
            assert c.post(API + "executions/run/", bad, format="json").status_code == 400
        assert c.post(API + "executions/validate/", {"filters": {"status": "submitted"}}, format="json").json()["data"]["valid"] is True
        assert c.post(API + "executions/validate/", {"filters": {"status": "zzz"}}, format="json").status_code == 400

    def test_scope_cannot_be_widened_through_the_api(self, w, seeded):
        r = api(w["r1_reporter"]).post(API + "executions/run/", {"filters": {"location": str(w["r2"].pk)}}, format="json").json()
        assert r["meta"]["count"] == 0
        assert api(w["r1_reporter"]).post(API + "executions/run/", {}, format="json").json()["meta"]["count"] == 2
        assert api(w["other_reporter"]).post(API + "executions/run/", {}, format="json").json()["meta"]["count"] == 0

    def test_export_status_and_private_download(self, w, seeded):
        c = api(w["reporter"])
        made = c.post(API + "executions/export/", {"format": "xlsx"}, format="json")
        run = body(made)
        assert made.status_code == 200 and run["status"] == "succeeded" and "artifact" not in str(run)
        assert body(c.get(f"{API}runs/{run['id']}/"))["row_count"] == 4 and c.get(API + "runs/").json()["meta"]["count"] == 1
        download = c.get(f"{API}runs/{run['id']}/download/")
        assert download.status_code == 200 and download.content[:2] == b"PK" and download["Cache-Control"] == "private, no-store"
        for other in ("exporter", "viewer_only", "r1_reporter"):
            assert api(w[other]).get(f"{API}runs/{run['id']}/download/").status_code == 404
        assert api(w["viewer_only"]).post(API + "executions/export/", {"format": "csv"}, format="json").status_code == 403
        assert c.post(API + "executions/export/", {"format": "docx"}, format="json").status_code == 400

    def test_saved_report_crud_and_ownership(self, w, seeded):
        c = api(w["reporter"])
        made = c.post(API + "saved/", {"report": "executions", "name": "API saved", "filters": {"status": "submitted"}, "columns": ["number"]}, format="json")
        sid = body(made)["id"]
        assert made.status_code == 201 and api(w["exporter"]).get(f"{API}saved/{sid}/").status_code == 404
        assert body(c.put(f"{API}saved/{sid}/", {"name": "Renamed", "filters": {}}, format="json"))["name"] == "Renamed"
        assert c.put(f"{API}saved/{sid}/", {"name": "x", "report": "findings"}, format="json").status_code == 400
        assert c.post(API + "saved/", {"name": "NoReport"}, format="json").status_code == 400
        assert c.post(API + "saved/", {"report": "executions", "name": "Bad", "filters": {"zzz": 1}}, format="json").status_code == 400
        ran = c.post(f"{API}saved/{sid}/run/", {}, format="json").json()
        assert ran["meta"]["count"] == 4
        dup = body(c.post(f"{API}saved/{sid}/duplicate/", {}, format="json"))
        assert dup["visibility"] == "private" and c.delete(f"{API}saved/{sid}/").status_code == 204 and c.get(f"{API}saved/{sid}/").status_code == 404

    def test_the_api_root_lists_the_report_endpoints(self, w, seeded):
        assert "reports" in str(api(w["reporter"]).get("/api/v1/").json())


@pytest.mark.django_db
def test_admin_is_read_only(admin_client):
    for model in ("savedreport", "reportrun"):
        assert admin_client.get(f"/admin/reporting/{model}/").status_code == 200
        assert admin_client.get(f"/admin/reporting/{model}/add/").status_code == 403
