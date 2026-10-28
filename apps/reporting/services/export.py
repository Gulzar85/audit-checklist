"""Export writers (Phase 19): CSV, XLSX (openpyxl) and PDF (a small built-in writer).  They only format rows the execution service produced - authorization and scope are
settled before a writer is called, so a column the user may not see never reaches a file.

Formula-injection policy (documented in docs/reporting-exports-and-artifacts.md): any *text* value that begins with ``=``, ``+``, ``-``, ``@``, tab or carriage return gets a leading
apostrophe in CSV and XLSX, so spreadsheet software shows it as text.  Numbers, booleans and dates are written as such and are never affected.
"""

from __future__ import annotations

import csv
import io
import re
import zlib
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

from django.utils import timezone

from .. import constants as k
from ..registry import BOOL, DATE, DATETIME, DECIMAL, INT, TEXT
from .execution import Result

DURATION = "duration"


def neutralize(text: str) -> str:
    return "'" + text if text.startswith(k.FORMULA_PREFIXES) else text


def _local(value: datetime) -> datetime:
    return timezone.localtime(value) if timezone.is_aware(value) else value


def text_of(value: Any, kind: str) -> str:
    """The display text of a value: ISO-style dates in the server time zone, plain numbers, Yes/No, minutes for durations."""
    if value is None:
        return ""
    if kind == DURATION and hasattr(value, "total_seconds"):
        return f"{value.total_seconds() / 60:.1f}"
    if isinstance(value, datetime):
        return _local(value).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (Decimal, float)):
        return f"{value:.2f}".rstrip("0").rstrip(".") if kind != INT else str(int(value))
    return str(value)


def filter_summary(result: Result) -> list[str]:
    lines = []
    report = result.report
    start, end = result.params.start, result.params.end
    if start and end:
        lines.append(f"Period: {_local(start):%Y-%m-%d} to {_local(end - timedelta(seconds=1)):%Y-%m-%d} ({report.date_label or 'all dates'})")
    for key, value in result.params.filters.items():
        if key not in ("date_from", "date_to"):
            lines.append(f"{key.replace('_', ' ').capitalize()}: {value}")
    if result.params.group_by:
        lines.append(f"Grouped by: {result.params.group_by}")
    return lines


# --------------------------------------------------------------------------------------------------------------------------------- CSV
def to_csv(result: Result) -> bytes:
    """UTF-8 with a byte-order mark (so Excel reads accented characters correctly), CRLF line ends, stable column order, empty cells for nulls."""
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow([neutralize(c.label) for c in result.columns])
    for row in result.rows:
        writer.writerow([neutralize(text_of(row.get(c.key), c.kind)) if c.kind == TEXT else text_of(row.get(c.key), c.kind) for c in result.columns])
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")


# -------------------------------------------------------------------------------------------------------------------------------- XLSX
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")


def to_xlsx(result: Result) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    book = Workbook(write_only=True)
    sheet = book.create_sheet(_BAD_SHEET.sub(" ", result.report.title)[:31] or "Report")
    bold = Font(bold=True)
    header = []
    for column in result.columns:
        cell = WriteOnlyCell(sheet, value=neutralize(column.label))
        cell.font = bold
        header.append(cell)
    sheet.append(header)
    for row in result.rows:
        cells: list[Any] = []
        for column in result.columns:
            value = row.get(column.key)
            cell = WriteOnlyCell(sheet)
            if value is None:
                cell.value = None
            elif column.kind == DURATION and hasattr(value, "total_seconds"):
                cell.value, cell.number_format = round(value.total_seconds() / 60, 1), "0.0"
            elif isinstance(value, bool):
                cell.value = value
            elif isinstance(value, datetime):
                cell.value, cell.number_format = _local(value).replace(tzinfo=None), "yyyy-mm-dd hh:mm"
            elif isinstance(value, date):
                cell.value, cell.number_format = value, "yyyy-mm-dd"
            elif isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
                cell.value = int(value) if column.kind == INT else float(value)
                cell.number_format = "0" if column.kind == INT else "0.00"
            else:
                cell.value = neutralize(str(value))
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)
    meta = book.create_sheet("Filters")
    meta.append(["Report", result.report.title])
    meta.append(["Version", result.report.version])
    meta.append(["Generated", _local(result.generated_at).replace(tzinfo=None)])
    for line in filter_summary(result):
        meta.append([neutralize(line)])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------------------------------------------------------------- PDF
PAGE_W, PAGE_H, MARGIN = 842, 595, 28  # A4 landscape, points
FONT = 7.5
LINE = 9.5


def _pdf_text(value: str) -> str:
    """Latin-1 (WinAnsi) text with PDF string escapes.  Characters outside Latin-1 become '?': the built-in writer ships no font files (documented limitation)."""
    encoded = value.encode("latin-1", "replace").decode("latin-1")
    return encoded.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _wrap(text: str, width_chars: int, max_lines: int = 3) -> list[str]:
    words, lines, current = text.replace("\n", " ").split(" "), [], ""
    for word in words:
        while len(word) > width_chars:  # a long unbroken token is split
            if current:
                lines.append(current)
                current = ""
            lines.append(word[:width_chars])
            word = word[width_chars:]
        candidate = f"{current} {word}".strip()
        if len(candidate) <= width_chars:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    lines = [ln for ln in lines if ln != ""] or [""]
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1][: max(0, width_chars - 3)] + "..."
    return lines


def to_pdf(result: Result) -> bytes:
    """A print-friendly table: title, generation time, filters and organization/period context on the first page; repeated column headers and 'Page x of y' on every page."""
    columns, rows = result.columns, result.rows
    usable = PAGE_W - 2 * MARGIN
    weights = []
    for c in columns:
        longest = max([len(c.label)] + [len(text_of(r.get(c.key), c.kind)) for r in rows[:200]])
        weights.append(min(max(longest, 6), 28))
    total_w = sum(weights)
    widths = [usable * w / total_w for w in weights]
    chars = [max(4, int(w / (FONT * 0.52)) - 1) for w in widths]
    pages: list[list[str]] = []
    ops: list[str] = []
    y = PAGE_H - MARGIN

    def start_page(first: bool) -> float:
        nonlocal ops
        ops = []
        top = PAGE_H - MARGIN
        if first:
            ops.append(f"BT /F2 14 Tf {MARGIN} {top - 12} Td ({_pdf_text(result.report.title)}) Tj ET")
            top -= 26
            meta = [f"Generated {_local(result.generated_at):%Y-%m-%d %H:%M} ({timezone.get_default_timezone_name()})", f"Organization: {getattr(result.organization, 'name', '')}", *filter_summary(result),
                    f"{result.total} row(s)"]
            for line in meta:
                ops.append(f"BT /F1 8 Tf {MARGIN} {top - 8} Td ({_pdf_text(line[:150])}) Tj ET")
                top -= 11
            top -= 6
        return table_header(top)

    def table_header(top: float) -> float:
        x = MARGIN
        ops.append(f"0.92 g {MARGIN} {top - 12} {usable} 14 re f 0 g")
        for c, w, n in zip(columns, widths, chars):
            ops.append(f"BT /F2 {FONT} Tf {x + 2} {top - 9} Td ({_pdf_text(c.label[:n])}) Tj ET")
            x += w
        return top - 16

    y = start_page(True)
    for row in rows:
        cells = [_wrap(text_of(row.get(c.key), c.kind), n) for c, n in zip(columns, chars)]
        height = max(len(c) for c in cells) * LINE + 2
        if y - height < MARGIN + 14:
            pages.append(ops)
            y = start_page(False)
        x = MARGIN
        for lines, w in zip(cells, widths):
            for i, line in enumerate(lines):
                ops.append(f"BT /F1 {FONT} Tf {x + 2} {y - 8 - i * LINE} Td ({_pdf_text(line)}) Tj ET")
            x += w
        ops.append(f"0.85 G {MARGIN} {y - height + 1} m {MARGIN + usable} {y - height + 1} l S 0 G")
        y -= height
    pages.append(ops)
    # assemble the file: 1 catalog, 2 pages, 3 Helvetica, 4 Helvetica-Bold, then per page a content stream and a page object
    objects: list[bytes] = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>"]
    kids = []
    count = len(pages)
    for index, page_ops in enumerate(pages, start=1):
        footer = f"BT /F1 7.5 Tf {PAGE_W - MARGIN - 60} {MARGIN - 12} Td (Page {index} of {count}) Tj ET"
        stream = zlib.compress("\n".join([*page_ops, footer]).encode("latin-1", "replace"))
        objects.append(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(stream) + stream + b"\nendstream")
        content_id = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {PAGE_W} {PAGE_H}] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content_id} 0 R >>".encode())
        kids.append(f"{len(objects)} 0 R")
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {count} >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info << /Title ({_pdf_text(result.report.title)}) >> >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


WRITERS = {k.Format.CSV: (to_csv, "text/csv; charset=utf-8", "csv"), k.Format.XLSX: (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
           k.Format.PDF: (to_pdf, "application/pdf", "pdf")}


def render(result: Result, fmt: str) -> tuple[bytes, str, str]:
    if fmt not in WRITERS or fmt not in result.report.formats:
        from ..errors import ReportError

        raise ReportError("That export format is not supported for this report.", "unsupported_format")
    writer, mime, ext = WRITERS[fmt]
    return writer(result), mime, ext
